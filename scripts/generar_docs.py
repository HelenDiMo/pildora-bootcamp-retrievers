"""Genera docs/teoria.docx y docs/teoria.pdf a partir de docs/teoria.md.

Hace falta tener instalados:
  - pandoc            (https://pandoc.org)
  - LibreOffice       (para convertir el .docx a PDF; comando `soffice`)
  - python-docx       (pip install python-docx; para dar bordes a las tablas)

Uso, desde la raíz del repositorio:

    python scripts/generar_docs.py
"""

import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

DOCS = Path(__file__).resolve().parents[1] / "docs"


def con_bordes(ruta_docx: Path) -> None:
    """Pandoc crea las tablas sin bordes; aquí se les pone borde y cabecera con fondo."""
    doc = Document(str(ruta_docx))
    for tabla in doc.tables:
        bordes = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{lado}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), "8A7A55")
            bordes.append(e)
        tabla._tbl.tblPr.append(bordes)
        for celda in tabla.rows[0].cells:
            sombra = OxmlElement("w:shd")
            sombra.set(qn("w:val"), "clear")
            sombra.set(qn("w:color"), "auto")
            sombra.set(qn("w:fill"), "E8DFC2")
            celda._tc.get_or_add_tcPr().append(sombra)
        for fila in tabla.rows:
            for celda in fila.cells:
                for parrafo in celda.paragraphs:
                    for run in parrafo.runs:
                        run.font.size = Pt(9.5)
    doc.save(str(ruta_docx))


def main() -> None:
    for programa in ("pandoc", "soffice"):
        if not shutil.which(programa):
            sys.exit(f"Falta '{programa}'. Instálalo y vuelve a intentarlo.")
    docx = DOCS / "teoria.docx"
    subprocess.run(
        ["pandoc", str(DOCS / "teoria.md"), "-o", str(docx), "--toc", "--toc-depth=2"], check=True
    )
    con_bordes(docx)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(DOCS), str(docx)],
        check=True,
        stdout=subprocess.DEVNULL,
    )
    print("Generados:", docx.name, "y", "teoria.pdf")


if __name__ == "__main__":
    main()
